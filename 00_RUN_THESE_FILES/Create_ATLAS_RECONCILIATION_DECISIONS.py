#!/usr/bin/env python3
"""
Create ATLAS_RECONCILIATION_DECISIONS.xlsx from Atlas reconciliation CSV outputs.

Put this file beside Run_NMR_ABET_Atlas.bat and NMR_ABET_Atlas.py.
The batch file runs this automatically after a successful Atlas scan.
"""

from __future__ import annotations

import csv
import re
import sys
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter
except ImportError:
    print("ERROR: openpyxl is required to create the Excel decision workbook.")
    print("Run Install_or_Update_Python_Packages.bat, then run Atlas again.")
    sys.exit(1)


OUTPUT_NAME = "ATLAS_RECONCILIATION_DECISIONS.xlsx"


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def newest_csv(folder: Path, pattern: str) -> Path | None:
    matches = list(folder.glob(pattern))
    if not matches:
        return None
    return max(matches, key=lambda p: p.stat().st_mtime)


def candidate_output_folders(script_dir: Path) -> list[Path]:
    folders = [
        script_dir,
        script_dir.parent / "03_ATLAS_OUTPUTS",
        script_dir.parent / "03_ATLAS_OUTPUTS" / "reports",
        script_dir / "03_ATLAS_OUTPUTS",
        script_dir / "03_ATLAS_OUTPUTS" / "reports",
        Path.cwd(),
        Path.cwd() / "reports",
    ]
    unique = []
    seen = set()
    for folder in folders:
        try:
            resolved = folder.resolve()
        except OSError:
            continue
        if resolved not in seen:
            seen.add(resolved)
            unique.append(folder)
    return unique


def find_atlas_outputs(script_dir: Path) -> tuple[Path, Path, Path | None, Path]:
    for folder in candidate_output_folders(script_dir):
        if not folder.exists():
            continue
        qc = newest_csv(folder, "ATLAS_QC_RECONCILIATION*.csv")
        coverage = newest_csv(folder, "ATLAS_ANIMAL_COVERAGE*.csv")
        metadata = newest_csv(folder, "ATLAS_ANIMAL_METADATA*.csv")
        if qc and coverage:
            return qc, coverage, metadata, folder

    print("ERROR: Could not find the Atlas reconciliation CSV outputs.")
    print("Looked for ATLAS_QC_RECONCILIATION*.csv and ATLAS_ANIMAL_COVERAGE*.csv in:")
    for folder in candidate_output_folders(script_dir):
        print(f"  - {folder}")
    sys.exit(1)


def is_standard_animal_id(animal_id: str) -> bool:
    return bool(re.match(r"^AS\d{2,4}[MF]\d+$", str(animal_id or "").strip(), flags=re.I))


def review_category(row: dict[str, str]) -> str:
    total = int(float(row.get("Total_Sessions") or 0))
    probe = int(float(row.get("Probe_2b") or 0))
    db_count = int(float(row.get("Databases_Count") or 0))
    animal_id = row.get("Animal_ID", "")

    if total >= 20 or probe >= 4 or db_count >= 2:
        return "HIGH PRIORITY - likely real study animal"
    if is_standard_animal_id(animal_id):
        return "Review - standard animal ID format"
    return "Likely exclude - nonstandard/test ID"


def suggested_action(row: dict[str, str]) -> str:
    animal_id = row.get("Animal_ID", "")
    if not is_standard_animal_id(animal_id):
        return "EXCLUDE unless you recognize this as a real animal"
    if review_category(row).startswith("HIGH PRIORITY"):
        return "KEEP or REMAP after checking metadata/source of truth"
    return "Review"


def default_decision(row: dict[str, str]) -> str:
    animal_id = row.get("Animal_ID", "")
    total = int(float(row.get("Total_Sessions") or 0))
    if not is_standard_animal_id(animal_id) and total <= 2:
        return "EXCLUDE"
    return ""


def autosize(ws) -> None:
    for col_idx in range(1, ws.max_column + 1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for cell in ws[col_letter]:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, min(len(value), 80))
        ws.column_dimensions[col_letter].width = max(12, min(max_len + 2, 55))


def style_sheet(ws, freeze: str = "A2") -> None:
    ws.freeze_panes = freeze
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.border = Border(bottom=thin)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.auto_filter.ref = ws.dimensions
    autosize(ws)


def add_validation(ws, col_letter: str, formula: str, last_row: int) -> None:
    if last_row < 2:
        return
    dv = DataValidation(type="list", formula1=f'"{formula}"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}2:{col_letter}{last_row}")


def write_rows(ws, headers: list[str], rows: list[dict[str, object]]) -> None:
    ws.append(headers)
    for row in rows:
        ws.append([row.get(header, "") for header in headers])


def build_workbook(qc_path: Path, coverage_path: Path, metadata_path: Path | None, output_folder: Path) -> Path:
    qc_rows = read_csv(qc_path)
    coverage_rows = read_csv(coverage_path)

    coverage_by_key = {
        (r.get("Project_Hint", ""), r.get("Animal_ID", "")): r
        for r in coverage_rows
    }

    unexpected_rows = []
    decision_rows = []
    cross_project_ids = {}

    for row in coverage_rows:
        animal_id = row.get("Animal_ID", "")
        cross_project_ids.setdefault(animal_id, []).append(row)

    flagged = []
    for qc in qc_rows:
        key = (qc.get("Project", ""), qc.get("Animal_ID", ""))
        merged = dict(coverage_by_key.get(key, {}))
        merged.update({
            "Review_Category": review_category(merged),
            "Review_Reason": qc.get("Detail", ""),
            "Suggested_Action": suggested_action(merged),
        })
        flagged.append(merged)

    flagged.sort(key=lambda r: (
        0 if str(r.get("Review_Category", "")).startswith("HIGH PRIORITY") else 1,
        r.get("Project_Hint", ""),
        r.get("Animal_ID", ""),
    ))

    for row in flagged:
        unexpected_rows.append(row)
        decision_rows.append({
            "Atlas_DB_Animal_ID": row.get("Animal_ID", ""),
            "Atlas_Project_Hint": row.get("Project_Hint", ""),
            "Decision": default_decision(row),
            "Correct_Animal_ID": "",
            "Final_Project_Group": row.get("Project_Hint", ""),
            "Final_Cohort": row.get("Cohort", ""),
            "Final_Sex": row.get("Sex", ""),
            "Final_Genotype": row.get("Genotype", ""),
            "Final_Treatment": row.get("Treatment", ""),
            "Reviewer_Notes": "",
            "Review_Category": row.get("Review_Category", ""),
            "Suggested_Action": row.get("Suggested_Action", ""),
            "Total_Sessions": row.get("Total_Sessions", ""),
            "Probe_2b": row.get("Probe_2b", ""),
            "Databases": row.get("Databases", ""),
            "First_Session": row.get("First_Session", ""),
            "Last_Session": row.get("Last_Session", ""),
            "Metadata_Match_Status": row.get("Metadata_Match_Status", ""),
            "Source_QC_File": qc_path.name,
            "Source_Coverage_File": coverage_path.name,
        })

    cross_rows = []
    for animal_id, rows in cross_project_ids.items():
        projects = {r.get("Project_Hint", "") for r in rows}
        if len(projects) > 1:
            cross_rows.extend(rows)
    cross_rows.sort(key=lambda r: (r.get("Animal_ID", ""), r.get("Project_Hint", "")))

    wb = Workbook()
    wb.remove(wb.active)

    how = wb.create_sheet("How To Use")
    how.append(["Step", "What to do", "Atlas behavior after patch"])
    how_rows = [
        [1, "Open the Decisions sheet.", "Atlas/extractor reads this file before clean analysis outputs."],
        [2, "Only review rows where Decision is blank or UNSURE.", "Rows still blank stay visible as unresolved review items."],
        [3, "Use KEEP when this is a real study animal and the ID is already correct.", "KEEP includes the animal and uses the final metadata you supplied."],
        [4, "Use REMAP when the database ID is a typo/alternate ID and enter the real ID in Correct_Animal_ID.", "REMAP replaces Atlas_DB_Animal_ID with Correct_Animal_ID before metadata matching."],
        [5, "Use EXCLUDE for test entries, wrong animals, junk IDs, or sessions that should never enter analysis.", "EXCLUDE removes those sessions/trials from clean analysis outputs but keeps them in QC audit logs."],
        [6, "Save the file as ATLAS_RECONCILIATION_DECISIONS.xlsx.", "Later scripts can find it by exact filename."],
    ]
    for row in how_rows:
        how.append(row)
    style_sheet(how)

    decisions = wb.create_sheet("Decisions")
    decision_headers = [
        "Atlas_DB_Animal_ID", "Atlas_Project_Hint", "Decision", "Correct_Animal_ID",
        "Final_Project_Group", "Final_Cohort", "Final_Sex", "Final_Genotype",
        "Final_Treatment", "Reviewer_Notes", "Review_Category", "Suggested_Action",
        "Total_Sessions", "Probe_2b", "Databases", "First_Session", "Last_Session",
        "Metadata_Match_Status", "Source_QC_File", "Source_Coverage_File",
    ]
    write_rows(decisions, decision_headers, decision_rows)
    style_sheet(decisions)
    last = decisions.max_row
    add_validation(decisions, "C", "KEEP,EXCLUDE,REMAP,UNSURE", last)
    add_validation(decisions, "E", "Lecanemab,Vaccine,Other,Unknown", last)
    add_validation(decisions, "G", "Male,Female,Unknown", last)
    add_validation(decisions, "H", "KI2TA-3,KI2TA-4,C57,AppKI,WT,Unknown", last)
    add_validation(decisions, "I", "PBS,Lecanemab,Sham,Vaccine,Unknown", last)

    cross = wb.create_sheet("Cross Project Check")
    cross_headers = [
        "Animal_ID", "Project_Hint", "Metadata_Match_Status", "Metadata_Project",
        "Cohort", "Sex", "Genotype", "Treatment", "Total_Sessions", "Probe_2b",
        "Databases", "First_Session", "Last_Session",
    ]
    write_rows(cross, cross_headers, cross_rows)
    style_sheet(cross)

    original = wb.create_sheet("Original Unexpected IDs")
    original_headers = [
        "Review_Category", "Project_Hint", "Animal_ID", "Review_Reason",
        "Suggested_Action", "Metadata_Match_Status", "Cohort", "Sex", "Genotype",
        "Treatment", "Total_Sessions", "Databases_Count", "Databases",
        "First_Session", "Last_Session", "Sessions_With_Trial_Data", "Stage_4",
        "Probe_2b", "CPT_Sessions",
    ]
    write_rows(original, original_headers, unexpected_rows)
    style_sheet(original)

    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = True

    output_path = output_folder / OUTPUT_NAME
    if output_path.exists():
        backup = output_folder / f"ATLAS_RECONCILIATION_DECISIONS_previous.xlsx"
        try:
            output_path.replace(backup)
            print(f"Existing {OUTPUT_NAME} moved to {backup.name}")
        except OSError:
            print(f"WARNING: Could not move existing {OUTPUT_NAME}; it may be open in Excel.")
            print("Close Excel and run this helper again.")
            sys.exit(1)

    wb.save(output_path)
    print(f"Created: {output_path}")
    print(f"Decision rows: {len(decision_rows)}")
    if metadata_path:
        print(f"Metadata source found: {metadata_path.name}")
    return output_path


def main() -> int:
    script_dir = Path(__file__).resolve().parent
    qc_path, coverage_path, metadata_path, output_folder = find_atlas_outputs(script_dir)
    print(f"Using QC file: {qc_path}")
    print(f"Using coverage file: {coverage_path}")
    build_workbook(qc_path, coverage_path, metadata_path, output_folder)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
