#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import os
import sys
import threading
import traceback
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Tuple

from atlas_core.indexer import build_atlas


def _norm_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _norm_sid(value) -> str:
    s = _norm_text(value)
    if not s:
        return ""
    try:
        f = float(s)
        if f.is_integer():
            return str(int(f))
    except Exception:
        pass
    return s


def _norm_db(value) -> str:
    s = _norm_text(value)
    return Path(s).name.lower() if s else ""


def _decision_key(source_database, sid) -> Tuple[str, str]:
    return (_norm_db(source_database), _norm_sid(sid))


def _read_probe_decision_rows(path: Path) -> List[dict]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as f:
            return list(csv.DictReader(f))

    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        sheet_name = "Decision_Queue" if "Decision_Queue" in wb.sheetnames else wb.sheetnames[0]
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [_norm_text(v) for v in rows[0]]
        out = []
        for vals in rows[1:]:
            if vals is None or not any(v is not None and _norm_text(v) for v in vals):
                continue
            out.append({headers[i]: vals[i] if i < len(vals) else None for i in range(len(headers))})
        return out

    raise ValueError(f"Unsupported Probe session decision file: {path}")


def _discover_probe_decisions(explicit: Optional[Path], output_dir: Path) -> Optional[Path]:
    if explicit:
        return explicit

    script_dir = Path(__file__).resolve().parent
    candidates = [
        output_dir / "ATLAS_PROBE_SESSION_DECISIONS.xlsx",
        output_dir / "ATLAS_PROBE_SESSION_DECISIONS.csv",
        script_dir / "ATLAS_PROBE_SESSION_DECISIONS.xlsx",
        script_dir / "ATLAS_PROBE_SESSION_DECISIONS.csv",
    ]
    for p in candidates:
        if p.exists():
            return p
    return None


def apply_probe_session_decisions(output_dir: Path, decision_path: Path, progress=print) -> dict:
    """Apply reviewed Probe 2b session-level QC decisions after reconciliation.

    This deliberately does not modify the raw Atlas cache or the reconciled
    ATLAS_SESSION_MASTER.csv. It writes a final analysis-facing session master
    plus an audit report.
    """
    reports = output_dir / "reports"
    clean_master = reports / "ATLAS_SESSION_MASTER.csv"
    if not clean_master.exists():
        raise FileNotFoundError(f"Cannot apply Probe session QC: {clean_master} was not created.")

    raw_rows = _read_probe_decision_rows(decision_path)
    if not raw_rows:
        raise ValueError(f"Probe session decision file contains no decision rows: {decision_path}")

    required = {"Source_Database", "SID", "Final_Decision"}
    missing = required - set(raw_rows[0].keys())
    if missing:
        raise ValueError(
            "Probe session decision file is missing required column(s): " + ", ".join(sorted(missing))
        )

    allowed = {"USE", "FAILED-REPLACED", "EXTRA"}
    decisions: Dict[Tuple[str, str], dict] = {}
    duplicate_keys = []
    unresolved = []

    for row in raw_rows:
        final = _norm_text(row.get("Final_Decision")).upper()
        if not final:
            unresolved.append(row)
            continue
        if final not in allowed:
            raise ValueError(
                f"Invalid Final_Decision {final!r} for {row.get('Animal_ID','')} / SID {row.get('SID','')}. "
                f"Allowed values: {', '.join(sorted(allowed))}."
            )
        key = _decision_key(row.get("Source_Database"), row.get("SID"))
        if not all(key):
            raise ValueError(
                f"Decision row is missing Source_Database or SID for animal {row.get('Animal_ID','')}."
            )
        if key in decisions:
            duplicate_keys.append(key)
        decisions[key] = {
            "decision": final,
            "reason": _norm_text(row.get("Reason")),
            "notes": _norm_text(row.get("Reviewer_Notes")),
            "project": _norm_text(row.get("Project")),
            "animal": _norm_text(row.get("Animal_ID")),
            "session_datetime": _norm_text(row.get("Session_DateTime")),
            "source_database": _norm_text(row.get("Source_Database")),
            "sid": _norm_sid(row.get("SID")),
        }

    if duplicate_keys:
        sample = ", ".join(f"{db}/SID {sid}" for db, sid in duplicate_keys[:5])
        raise ValueError(f"Duplicate Probe QC decision key(s) found: {sample}")
    if unresolved:
        raise ValueError(
            f"Probe decision file still has {len(unresolved)} blank Final_Decision row(s). "
            "Finish the Decision_Queue before running Atlas."
        )

    final_master = reports / "ATLAS_SESSION_MASTER_FINAL.csv"
    audit_path = reports / "ATLAS_PROBE_SESSION_QC_APPLIED.csv"

    matched = set()
    included = excluded_failed = excluded_extra = 0

    with clean_master.open("r", encoding="utf-8-sig", newline="") as src:
        reader = csv.DictReader(src)
        fieldnames = list(reader.fieldnames or [])
        added = [
            "Probe_QC_Decision",
            "Probe_QC_Reason",
            "Probe_QC_Reviewer_Notes",
            "Include_Primary_Analysis",
        ]
        out_fields = fieldnames + [x for x in added if x not in fieldnames]

        with final_master.open("w", encoding="utf-8-sig", newline="") as dst:
            writer = csv.DictWriter(dst, fieldnames=out_fields)
            writer.writeheader()
            for row in reader:
                key = _decision_key(row.get("Source_Database"), row.get("SID"))
                dec = decisions.get(key)
                if dec:
                    # Guard against accidental matching to the wrong animal/project.
                    if dec["animal"] and _norm_text(row.get("Animal_ID")) != dec["animal"]:
                        continue
                    if dec["project"] and _norm_text(row.get("Project_Hint")) != dec["project"]:
                        continue
                    matched.add(key)
                    row["Probe_QC_Decision"] = dec["decision"]
                    row["Probe_QC_Reason"] = dec["reason"]
                    row["Probe_QC_Reviewer_Notes"] = dec["notes"]
                    if dec["decision"] == "USE":
                        row["Include_Primary_Analysis"] = "YES"
                        writer.writerow(row)
                        included += 1
                    elif dec["decision"] == "FAILED-REPLACED":
                        excluded_failed += 1
                    elif dec["decision"] == "EXTRA":
                        excluded_extra += 1
                else:
                    row["Probe_QC_Decision"] = ""
                    row["Probe_QC_Reason"] = ""
                    row["Probe_QC_Reviewer_Notes"] = ""
                    row["Include_Primary_Analysis"] = "YES"
                    writer.writerow(row)
                    included += 1

    unmatched = [k for k in decisions if k not in matched]

    audit_fields = [
        "Source_Database", "SID", "Project", "Animal_ID", "Session_DateTime",
        "Final_Decision", "Reason", "Reviewer_Notes", "Matched_In_Clean_Master",
        "Primary_Analysis_Action",
    ]
    with audit_path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=audit_fields)
        w.writeheader()
        for key, dec in decisions.items():
            decision = dec["decision"]
            w.writerow({
                "Source_Database": dec["source_database"],
                "SID": dec["sid"],
                "Project": dec["project"],
                "Animal_ID": dec["animal"],
                "Session_DateTime": dec["session_datetime"],
                "Final_Decision": decision,
                "Reason": dec["reason"],
                "Reviewer_Notes": dec["notes"],
                "Matched_In_Clean_Master": "YES" if key in matched else "NO",
                "Primary_Analysis_Action": "INCLUDE" if decision == "USE" else "EXCLUDE",
            })

    if unmatched:
        sample = ", ".join(f"{db}/SID {sid}" for db, sid in unmatched[:8])
        raise ValueError(
            f"Probe QC decisions were not fully applied: {len(unmatched)} decision row(s) did not match "
            f"ATLAS_SESSION_MASTER.csv. Example(s): {sample}. "
            f"Audit written to {audit_path}."
        )

    progress(
        f"Probe session QC applied: {len(decisions)} reviewed sessions; "
        f"{excluded_failed} FAILED-REPLACED excluded, {excluded_extra} EXTRA excluded."
    )
    progress(f"Final analysis session master: {final_master}")
    progress(f"Probe QC audit: {audit_path}")

    return {
        "final_session_master": final_master,
        "probe_qc_audit": audit_path,
        "reviewed": len(decisions),
        "failed_replaced": excluded_failed,
        "extra": excluded_extra,
        "included_rows": included,
    }


def run_cli(args: argparse.Namespace) -> int:
    def log(msg: str):
        print(msg, flush=True)

    try:
        output_dir = Path(args.output)
        result = build_atlas(
            Path(args.input),
            output_dir,
            Path(args.metadata) if args.metadata else None,
            workers=args.workers,
            force_rebuild=args.force_rebuild,
            export_trials_csv=args.export_trials_csv,
            progress=log,
            reconciliation_path=Path(args.reconciliation) if args.reconciliation else None,
        )

        probe_path = _discover_probe_decisions(Path(args.probe_decisions) if args.probe_decisions else None, output_dir)
        probe_result = None
        if probe_path:
            log(f"Applying Probe session QC decisions: {probe_path}")
            probe_result = apply_probe_session_decisions(output_dir, probe_path, progress=log)
        else:
            log("No Probe session QC decision file selected/found; ATLAS_SESSION_MASTER_FINAL.csv was not created.")

        print("\nDONE")
        print("Atlas cache:", result["index"])
        for name, path in result["outputs"].items():
            if str(path):
                print(f"{name}: {path}")
        if probe_result:
            print("final_session_master:", probe_result["final_session_master"])
            print("probe_qc_audit:", probe_result["probe_qc_audit"])
        return 0
    except Exception:
        traceback.print_exc()
        return 1


def run_gui(default_workers: int = 6) -> int:
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk

    root = tk.Tk()
    root.title("NMR ABET Atlas v1.1")
    root.geometry("920x720")
    root.minsize(820, 620)

    in_var = tk.StringVar()
    out_var = tk.StringVar()
    meta_var = tk.StringVar()
    recon_var = tk.StringVar()
    probe_var = tk.StringVar()
    workers_var = tk.StringVar(value=str(default_workers))
    trial_csv_var = tk.BooleanVar(value=False)
    force_var = tk.BooleanVar(value=False)

    frm = ttk.Frame(root, padding=14)
    frm.pack(fill="both", expand=True)
    ttk.Label(frm, text="NMR ABET Atlas", font=("Segoe UI", 20, "bold")).grid(row=0, column=0, columnspan=3, sticky="w")
    ttk.Label(
        frm,
        text="Index every animal/session once, apply reconciliation, then apply reviewed Probe session QC decisions.",
    ).grid(row=1, column=0, columnspan=3, sticky="w", pady=(0, 14))

    def browse_dir(var):
        p = filedialog.askdirectory()
        if p:
            var.set(p)

    def browse_meta():
        p = filedialog.askopenfilename(filetypes=[("CSV metadata", "*.csv"), ("All files", "*.*")])
        if p:
            meta_var.set(p)

    def browse_recon():
        p = filedialog.askopenfilename(
            filetypes=[
                ("Atlas reconciliation decisions", "ATLAS_RECONCILIATION_DECISIONS*.xlsx ATLAS_RECONCILIATION_DECISIONS*.csv"),
                ("Excel files", "*.xlsx"),
                ("CSV files", "*.csv"),
                ("All files", "*.*"),
            ]
        )
        if p:
            recon_var.set(p)

    def browse_probe():
        p = filedialog.askopenfilename(
            filetypes=[
                ("Atlas Probe session decisions", "ATLAS_PROBE_SESSION_DECISIONS*.xlsx ATLAS_PROBE_SESSION_DECISIONS*.csv"),
                ("Excel files", "*.xlsx"),
                ("CSV files", "*.csv"),
                ("All files", "*.*"),
            ]
        )
        if p:
            probe_var.set(p)

    labels = [
        ("ABET database folder", in_var, browse_dir),
        ("Atlas output folder", out_var, browse_dir),
    ]
    row = 2
    for label_text, var, fn in labels:
        ttk.Label(frm, text=label_text).grid(row=row, column=0, sticky="w", pady=5)
        ttk.Entry(frm, textvariable=var, width=72).grid(row=row, column=1, sticky="ew", padx=8)
        ttk.Button(frm, text="Browse…", command=lambda v=var, f=fn: f(v)).grid(row=row, column=2)
        row += 1

    ttk.Label(frm, text="Animal metadata CSV (optional)").grid(row=row, column=0, sticky="w", pady=5)
    ttk.Entry(frm, textvariable=meta_var, width=72).grid(row=row, column=1, sticky="ew", padx=8)
    ttk.Button(frm, text="Browse…", command=browse_meta).grid(row=row, column=2)
    row += 1

    ttk.Label(frm, text="Reconciliation decisions XLSX/CSV (optional)").grid(row=row, column=0, sticky="w", pady=5)
    ttk.Entry(frm, textvariable=recon_var, width=72).grid(row=row, column=1, sticky="ew", padx=8)
    ttk.Button(frm, text="Browse…", command=browse_recon).grid(row=row, column=2)
    row += 1

    ttk.Label(frm, text="Probe session QC decisions XLSX/CSV (optional)").grid(row=row, column=0, sticky="w", pady=5)
    ttk.Entry(frm, textvariable=probe_var, width=72).grid(row=row, column=1, sticky="ew", padx=8)
    ttk.Button(frm, text="Browse…", command=browse_probe).grid(row=row, column=2)
    row += 1
    ttk.Label(
        frm,
        text="If blank, Atlas auto-detects ATLAS_PROBE_SESSION_DECISIONS.xlsx/csv in the output folder or Atlas program folder.",
        foreground="#555555",
    ).grid(row=row, column=1, columnspan=2, sticky="w", padx=8, pady=(0, 6))
    row += 1

    ttk.Label(frm, text="Workers").grid(row=row, column=0, sticky="w", pady=5)
    ttk.Combobox(frm, textvariable=workers_var, values=["2", "4", "6", "8", "10", "12"], width=8, state="readonly").grid(row=row, column=1, sticky="w", padx=8)
    row += 1
    ttk.Checkbutton(
        frm,
        text="Export giant ATLAS_TRIAL_MASTER.csv too (optional; trials are always stored in Atlas SQLite cache)",
        variable=trial_csv_var,
    ).grid(row=row, column=0, columnspan=3, sticky="w", pady=3)
    row += 1
    ttk.Checkbutton(frm, text="Force rebuild all database caches", variable=force_var).grid(row=row, column=0, columnspan=3, sticky="w", pady=3)
    row += 1

    ttk.Separator(frm).grid(row=row, column=0, columnspan=3, sticky="ew", pady=10)
    row += 1
    ttk.Label(frm, text="Progress / audit log").grid(row=row, column=0, columnspan=3, sticky="w")
    row += 1
    text = tk.Text(frm, height=20, wrap="word", font=("Consolas", 9))
    text.grid(row=row, column=0, columnspan=3, sticky="nsew")
    row += 1
    sb = ttk.Scrollbar(frm, orient="vertical", command=text.yview)
    text.configure(yscrollcommand=sb.set)

    frm.columnconfigure(1, weight=1)
    frm.rowconfigure(row - 1, weight=1)

    def append(msg):
        text.insert("end", msg + "\n")
        text.see("end")
        root.update_idletasks()

    run_btn = None

    def start():
        nonlocal run_btn
        if not in_var.get() or not out_var.get():
            messagebox.showerror("Missing folder", "Choose both the ABET database folder and Atlas output folder.")
            return
        try:
            workers = int(workers_var.get())
        except Exception:
            workers = 6
        run_btn.config(state="disabled")
        text.delete("1.0", "end")
        append("Starting NMR ABET Atlas…")

        def worker():
            try:
                output_dir = Path(out_var.get())
                result = build_atlas(
                    Path(in_var.get()),
                    output_dir,
                    Path(meta_var.get()) if meta_var.get() else None,
                    workers=workers,
                    force_rebuild=force_var.get(),
                    export_trials_csv=trial_csv_var.get(),
                    progress=lambda m: root.after(0, append, m),
                    reconciliation_path=Path(recon_var.get()) if recon_var.get() else None,
                )

                probe_path = _discover_probe_decisions(Path(probe_var.get()) if probe_var.get() else None, output_dir)
                probe_result = None
                if probe_path:
                    root.after(0, append, f"Applying Probe session QC decisions: {probe_path}")
                    probe_result = apply_probe_session_decisions(output_dir, probe_path, progress=lambda m: root.after(0, append, m))
                else:
                    root.after(0, append, "No Probe session QC decision file selected/found; final analysis master not created.")

                msg = (
                    f"Done.\n\nSessions: {result['counts']['sessions']:,}\n"
                    f"Animal IDs: {result['counts']['animals']:,}\n"
                    f"Trials cached: {result['counts']['trials']:,}"
                )
                if probe_result:
                    msg += (
                        f"\n\nProbe QC applied: {probe_result['reviewed']} reviewed sessions"
                        f"\nFAILED-REPLACED excluded: {probe_result['failed_replaced']}"
                        f"\nEXTRA excluded: {probe_result['extra']}"
                        f"\n\nFinal master: ATLAS_SESSION_MASTER_FINAL.csv"
                    )
                msg += "\n\nReports are in the reports folder."
                root.after(0, lambda: messagebox.showinfo("Atlas complete", msg))
            except Exception as exc:
                detail = traceback.format_exc()
                root.after(0, append, detail)
                root.after(0, lambda: messagebox.showerror("Atlas error", str(exc)))
            finally:
                root.after(0, lambda: run_btn.config(state="normal"))

        threading.Thread(target=worker, daemon=True).start()

    run_btn = ttk.Button(frm, text="BUILD / UPDATE ATLAS", command=start)
    run_btn.grid(row=row, column=0, columnspan=3, pady=(12, 0), ipady=7, sticky="ew")
    root.mainloop()
    return 0


def main() -> int:
    p = argparse.ArgumentParser(description="NMR ABET Atlas — universal ABET session/trial indexer")
    p.add_argument("--input", help="Folder containing ABETdb/SQLite databases")
    p.add_argument("--output", help="Atlas output folder")
    p.add_argument("--metadata", help="Optional animal metadata CSV")
    p.add_argument("--reconciliation", help="Optional ATLAS_RECONCILIATION_DECISIONS.xlsx/csv file")
    p.add_argument("--probe-decisions", help="Optional ATLAS_PROBE_SESSION_DECISIONS.xlsx/csv file")
    p.add_argument("--workers", type=int, default=6)
    p.add_argument("--force-rebuild", action="store_true")
    p.add_argument("--export-trials-csv", action="store_true")
    p.add_argument("--no-gui", action="store_true")
    args = p.parse_args()
    if args.input and args.output:
        return run_cli(args)
    if args.no_gui:
        p.error("--input and --output are required with --no-gui")
    return run_gui(args.workers)


if __name__ == "__main__":
    raise SystemExit(main())
