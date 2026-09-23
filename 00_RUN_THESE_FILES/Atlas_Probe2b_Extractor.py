#!/usr/bin/env python3
from __future__ import annotations

import argparse, csv, math, sqlite3, statistics, sys, traceback
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from statistics import NormalDist
from typing import Any, Iterable

NORMAL = NormalDist()


# Exact legacy Taylor/Lecanemab 49-column project schema.
# Spelling, punctuation, capitalization and order are intentionally preserved.
TAYLOR_49_COLUMNS = [
    "AnimalID", "Age", "Sex", "Treatment", "Genotype", "Strain", "ExpName", "Housing",
    "LightCycle", "PISiteUser", "SessionName", "Image", "Image_Description", "sessionInfoName",
    "Intervention", "S_ID",
    "AVG_Average - Correct Choice Latency at 0.2s SD",
    "AVG_Average - Correct Choice Latency at 0.5s SD",
    "AVG_Average - Correct Choice Latency at 1s SD",
    "AVG_Average - Correct Choice Latency at 2s SD",
    "AVG_Average - Mistake Latency at 0.2s SD",
    "AVG_Average - Mistake Latency at 0.5s SD",
    "AVG_Average - Mistake Latency at 1s SD",
    "AVG_Average - Mistake Latency at 2s SD",
    "AVG_Average - Reward Retrieval Latency at 0.2s SD",
    "AVG_Average - Reward Retrieval Latency at 0.5s SD",
    "AVG_Average - Reward Retrieval Latency at 1s SD",
    "AVG_Average - Reward Retrieval Latency at 2s SD",
    "AVG_count at 0.2s", "AVG_count at 0.5s", "AVG_count at 1s", "AVG_count at 2s",
    "AVG_End Summary - CorrectRejection at 0.2s SD",
    "AVG_End Summary - CorrectRejection at 0.5s SD",
    "AVG_End Summary - CorrectRejection at 1s SD",
    "AVG_End Summary - CorrectRejection at 2s SD",
    "AVG_End Summary - Hits",
    "AVG_End Summary - Hits at 0.2s SD",
    "AVG_End Summary - Hits at 0.5s SD",
    "AVG_End Summary - Hits at 1s SD",
    "AVG_End Summary - Hits at 2s SD",
    "AVG_End Summary - Miss at 0.2s SD",
    "AVG_End Summary - Miss at 0.5s SD",
    "AVG_End Summary - Miss at 1s SD",
    "AVG_End Summary - Miss at 2s SD",
    "AVG_End Summary - Mistake at 0.2s SD",
    "AVG_End Summary - Mistake at 0.5s SD",
    "AVG_End Summary - Mistake at 1s SD",
    "AVG_End Summary - Mistake at 2s SD",
]
TAYLOR_DURATIONS = [(0.2,"0.2s"),(0.5,"0.5s"),(1.0,"1s"),(2.0,"2s")]


# Exact legacy 67-column extended CPT layout: Taylor 49 columns + Date_Time +
# chronological day + four SDT metrics for each of the four Probe 2b durations.
SDT_67_COLUMNS = []
for _label in ["0.2s", "0.5s", "1s", "2s"]:
    SDT_67_COLUMNS.extend([
        f"Hit Rate at {_label} SD",
        f"False Alarm Rate at {_label} SD",
        f"dPrime at {_label} SD",
        f"c Response Bias at {_label} SD",
    ])
LEGACY_67_COLUMNS = TAYLOR_49_COLUMNS + ["Date_Time", "chronological day"] + SDT_67_COLUMNS

# ----------------------------- helpers -----------------------------
def txt(x: Any) -> str:
    return "" if x is None else str(x).strip()

def num(x: Any):
    try:
        if x is None or x == "": return None
        return float(x)
    except Exception:
        return None

def parse_dt(s: Any):
    s=txt(s)
    if not s: return None
    for f in (None, "%Y-%m-%dT%H:%M:%S.%f", "%Y-%m-%dT%H:%M:%S", "%m/%d/%y %H:%M:%S"):
        try:
            return datetime.fromisoformat(s) if f is None else datetime.strptime(s,f)
        except Exception: pass
    return None

def read_csv(path: Path):
    with path.open(newline="", encoding="utf-8-sig", errors="replace") as f:
        return list(csv.DictReader(f))

def write_csv(path: Path, rows: list[dict[str,Any]], fields: list[str] | None=None):
    path.parent.mkdir(parents=True, exist_ok=True)
    if fields is None:
        fields=[]
        seen=set()
        for r in rows:
            for k in r:
                if k not in seen: seen.add(k); fields.append(k)
    with path.open("w",newline="",encoding="utf-8-sig") as f:
        w=csv.DictWriter(f,fieldnames=fields,extrasaction="ignore")
        w.writeheader(); w.writerows(rows)

def percentile_linear(vals: list[float], p: float) -> float:
    s=sorted(vals); n=len(s)
    if n==1: return s[0]
    pos=(n-1)*p; lo=math.floor(pos); hi=math.ceil(pos)
    if lo==hi: return s[lo]
    return s[lo]+(s[hi]-s[lo])*(pos-lo)

def iqr_bounds(vals: list[float]):
    if len(vals)<4: return None,None
    q1=percentile_linear(vals,0.25); q3=percentile_linear(vals,0.75); i=q3-q1
    return q1-1.5*i, q3+1.5*i

def rate_metrics(h:int,m:int,fa:int,cr:int):
    ns=h+m; nn=fa+cr
    hr=h/ns if ns else None; far=fa/nn if nn else None
    out={"Hit_Rate":hr,"False_Alarm_Rate":far,"Hit_Rate_for_dprime":None,"False_Alarm_Rate_for_dprime":None,"d_prime":None,"criterion_c":None}
    if hr is None or far is None: return out
    hrc=(h-0.5)/ns if hr==1 else (h+0.5)/ns if hr==0 else hr
    farc=(fa-0.5)/nn if far==1 else (fa+0.5)/nn if far==0 else far
    out["Hit_Rate_for_dprime"]=hrc; out["False_Alarm_Rate_for_dprime"]=farc
    if 0<hrc<1 and 0<farc<1:
        zh=NORMAL.inv_cdf(hrc); zf=NORMAL.inv_cdf(farc)
        out["d_prime"]=zh-zf; out["criterion_c"]=-(zh+zf)/2
    return out

def mean(vals: Iterable[Any]):
    v=[float(x) for x in vals if x is not None and x!="" and not (isinstance(x,float) and math.isnan(x))]
    return sum(v)/len(v) if v else None

def cluster_sessions(rows: list[dict[str,Any]], max_gap_days: float=7.0):
    rows=sorted(rows,key=lambda r: parse_dt(r.get("Session_DateTime")) or datetime.min)
    clusters=[]; cur=[]; prev=None
    for r in rows:
        dt=parse_dt(r.get("Session_DateTime"))
        if prev is None or dt is None or (dt-prev).total_seconds()/86400 <= max_gap_days:
            cur.append(r)
        else:
            clusters.append(cur); cur=[r]
        prev=dt
    if cur: clusters.append(cur)
    return clusters

def session_key(r):
    return (txt(r.get("Database_Fingerprint")), txt(r.get("SID")))

def sid_text(value: Any) -> str:
    """Normalize Atlas/MouseBytes session IDs for exact whitelist matching."""
    value = txt(value)
    if value.endswith(".0"):
        value = value[:-2]
    return value

def project_text(row: dict[str, Any]) -> str:
    return txt(row.get("Project_Group") or row.get("Project_Hint") or row.get("Project"))

def animal_text(row: dict[str, Any]) -> str:
    return txt(row.get("AnimalID") or row.get("Animal_ID") or row.get("Animal ID"))

def selection_key(row: dict[str, Any]) -> tuple[str, str, str]:
    # The reviewed spreadsheet S_ID and Atlas master SID use different
    # numbering systems. The corrected whitelist therefore carries the Atlas
    # SID explicitly; fall back to the original field for legacy files.
    sid = row.get("Atlas_Master_SID") or (row.get("Source_SID") if "Source_SID" in row else row.get("SID"))
    fingerprint = row.get("Atlas_Master_Database_Fingerprint") or row.get("Database_Fingerprint")
    sid_token = f"{txt(fingerprint)}::{sid_text(sid)}" if txt(fingerprint) else sid_text(sid)
    return (project_text(row).casefold(), animal_text(row).casefold(), sid_token)

def is_yes(v): return txt(v).upper() in {"YES","Y","TRUE","1"}

def norm_sd(v):
    x=num(v)
    if x is None: return None
    # Atlas stores seconds; tolerate milliseconds just in case.
    if x>20: x=x/1000.0
    for t in (2.0,1.0,0.5,0.2):
        if abs(x-t)<0.03: return t
    return round(x,4)



def find_metadata_csv(reports: Path) -> Path | None:
    """Locate the Atlas animal metadata table beside the session reports.

    Prefer the canonical filename; tolerate duplicate-safe Windows/download suffixes
    so a copied reports folder still works.
    """
    canonical = reports / "ATLAS_ANIMAL_METADATA.csv"
    if canonical.exists():
        return canonical
    candidates = sorted(reports.glob("ATLAS_ANIMAL_METADATA*.csv"), key=lambda p: p.stat().st_mtime, reverse=True)
    return candidates[0] if candidates else None

def find_selection_csv(reports: Path) -> Path | None:
    """Find an optional exact Probe 2b session whitelist beside Atlas reports."""
    canonical = reports / "ATLAS_PROBE2B_SESSION_SELECTIONS.csv"
    if canonical.exists():
        return canonical
    candidates = sorted(
        reports.glob("ATLAS_PROBE2B_SESSION_SELECTIONS*.csv"),
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )
    return candidates[0] if candidates else None

def find_reconciliation_workbook(atlas_dir: Path, reports: Path) -> Path | None:
    """Find the human-reviewed KEEP/EXCLUDE/REMAP decision workbook."""
    candidates = [
        reports / "ATLAS_RECONCILIATION_DECISIONS.xlsx",
        atlas_dir / "ATLAS_RECONCILIATION_DECISIONS.xlsx",
        atlas_dir / "reports" / "ATLAS_RECONCILIATION_DECISIONS.xlsx",
    ]
    for path in candidates:
        if path.exists():
            return path
    hits = []
    for folder in {reports, atlas_dir, atlas_dir / "reports"}:
        if folder.exists():
            hits.extend(folder.glob("ATLAS_RECONCILIATION_DECISIONS*.xlsx"))
    if not hits:
        return None
    return max(hits, key=lambda p: p.stat().st_mtime)

def read_decision_workbook(path: Path) -> list[dict[str, Any]]:
    """Read the Decisions sheet from ATLAS_RECONCILIATION_DECISIONS.xlsx."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError(
            "openpyxl is required to read ATLAS_RECONCILIATION_DECISIONS.xlsx. "
            "Run Install_or_Update_Python_Packages.bat, then run the extractor again."
        ) from exc

    wb = load_workbook(path, data_only=True, read_only=True)
    sheet_name = "Decisions" if "Decisions" in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet_name]
    headers = [txt(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        row = {headers[i]: values[i] for i in range(min(len(headers), len(values))) if headers[i]}
        if txt(row.get("Atlas_DB_Animal_ID")) or txt(row.get("Animal_ID")):
            rows.append(row)
    return rows

def decision_key(project: str, animal: str) -> tuple[str, str]:
    return (txt(project).casefold(), txt(animal).casefold())

def build_decision_maps(decision_rows: list[dict[str, Any]]):
    maps = {
        "keep": {},
        "exclude": {},
        "remap": {},
        "unresolved": {},
        "all": {},
    }
    for row in decision_rows:
        animal = txt(row.get("Atlas_DB_Animal_ID") or row.get("Animal_ID"))
        project = txt(row.get("Atlas_Project_Hint") or row.get("Final_Project_Group") or row.get("Project_Hint") or row.get("Project"))
        if not animal or not project:
            continue
        decision = txt(row.get("Decision")).upper()
        key = decision_key(project, animal)
        maps["all"][key] = row
        if decision == "KEEP":
            maps["keep"][key] = row
        elif decision == "EXCLUDE":
            maps["exclude"][key] = row
        elif decision == "REMAP":
            maps["remap"][key] = row
        else:
            # Blank and UNSURE are intentionally not treated as clean analysis animals.
            maps["unresolved"][key] = row
    return maps

def apply_final_decision_metadata(row: dict[str, Any], decision_row: dict[str, Any]) -> None:
    mapping = {
        "Final_Project_Group": ("Project_Hint", "Project_Group", "Project"),
        "Final_Cohort": ("Cohort",),
        "Final_Sex": ("Sex",),
        "Final_Genotype": ("Genotype", "Strain"),
        "Final_Treatment": ("Treatment",),
    }
    for source, targets in mapping.items():
        value = txt(decision_row.get(source))
        if not value:
            continue
        for target in targets:
            row[target] = value
    row["Metadata_Match_Status"] = "MATCHED_BY_RECONCILIATION"

def apply_reconciliation_decisions(master_rows: list[dict[str, Any]], decision_rows: list[dict[str, Any]]):
    """Apply KEEP/EXCLUDE/REMAP decisions before analysis extraction.

    The Atlas session master remains untouched on disk. This function returns
    patched in-memory rows plus an audit table that is written beside outputs.
    """
    maps = build_decision_maps(decision_rows)
    remap_targets = {
        key: txt(row.get("Correct_Animal_ID"))
        for key, row in maps["remap"].items()
        if txt(row.get("Correct_Animal_ID"))
    }
    output = []
    audit = []
    counts = defaultdict(int)

    for row0 in master_rows:
        row = row0.copy()
        project = txt(row.get("Project_Hint") or row.get("Project_Group") or row.get("Project"))
        animal = txt(row.get("Animal_ID") or row.get("AnimalID") or row.get("Animal ID"))
        key = decision_key(project, animal)
        action = "NO_DECISION_ROW"
        include = True
        reason = ""
        decision_row = None

        if key in maps["exclude"]:
            decision_row = maps["exclude"][key]
            action = "EXCLUDE"
            include = False
            reason = "Excluded by ATLAS_RECONCILIATION_DECISIONS.xlsx"
        elif key in maps["remap"]:
            decision_row = maps["remap"][key]
            corrected = remap_targets.get(key)
            if corrected:
                action = "REMAP"
                row["Original_Animal_ID"] = animal
                row["Animal_ID"] = corrected
                row["AnimalID"] = corrected
                apply_final_decision_metadata(row, decision_row)
                reason = f"Remapped {animal} to {corrected}"
            else:
                action = "REMAP_MISSING_CORRECT_ID"
                include = False
                reason = "REMAP selected but Correct_Animal_ID is blank"
        elif key in maps["keep"]:
            decision_row = maps["keep"][key]
            action = "KEEP"
            apply_final_decision_metadata(row, decision_row)
            reason = "Kept by reconciliation decision"
        elif key in maps["unresolved"]:
            decision_row = maps["unresolved"][key]
            action = "UNRESOLVED"
            include = False
            reason = "Decision is blank or UNSURE"

        row["Reconciliation_Action"] = action
        row["Reconciliation_Include"] = "YES" if include else "NO"
        row["Reconciliation_Reason"] = reason
        if include:
            output.append(row)
        counts[action] += 1
        if action != "NO_DECISION_ROW":
            audit.append({
                "Project": project,
                "Original_Animal_ID": animal,
                "Final_Animal_ID": txt(row.get("Animal_ID") or row.get("AnimalID")),
                "Action": action,
                "Include_In_Analysis": "YES" if include else "NO",
                "Reason": reason,
                "Final_Project_Group": txt(decision_row.get("Final_Project_Group")) if decision_row else "",
                "Final_Cohort": txt(decision_row.get("Final_Cohort")) if decision_row else "",
                "Final_Sex": txt(decision_row.get("Final_Sex")) if decision_row else "",
                "Final_Genotype": txt(decision_row.get("Final_Genotype")) if decision_row else "",
                "Final_Treatment": txt(decision_row.get("Final_Treatment")) if decision_row else "",
                "Reviewer_Notes": txt(decision_row.get("Reviewer_Notes")) if decision_row else "",
            })

    return output, audit, dict(counts)


def merge_animal_metadata(master_rows: list[dict[str,Any]], metadata_rows: list[dict[str,Any]]):
    """Attach finalized animal metadata using Project + Animal ID.

    This is deliberately project-aware because the same Animal ID may represent
    different animals in the Vaccine and Lecanemab projects. Existing nonblank
    session values are preserved; finalized animal metadata fills/overrides the
    group fields used for analysis and compatibility exports.
    """
    lookup={}
    for m in metadata_rows:
        project=txt(m.get("Project_Group") or m.get("Project_Hint") or m.get("Project"))
        animal=txt(m.get("AnimalID") or m.get("Animal_ID") or m.get("Animal ID"))
        if project and animal:
            lookup[(project.casefold(),animal.casefold())]=m
    fields=("Cohort","Sex","Genotype","Treatment","Strain")
    out=[]; matched=0
    for r0 in master_rows:
        r=r0.copy()
        project=txt(r.get("Project_Hint") or r.get("Project_Group") or r.get("Project"))
        animal=txt(r.get("Animal_ID") or r.get("AnimalID") or r.get("Animal ID"))
        m=lookup.get((project.casefold(),animal.casefold()))
        if m is not None:
            matched += 1
            for f in fields:
                mv=txt(m.get(f))
                if mv:
                    r[f]=mv
            # Strain is intentionally represented by genotype for these cohorts
            # when the metadata source does not carry a separate strain value.
            if not txt(r.get("Strain")) and txt(r.get("Genotype")):
                r["Strain"]=txt(r.get("Genotype"))
            r["Metadata_Match_Status"]="MATCHED"
        else:
            r["Metadata_Match_Status"]="UNMATCHED"
        out.append(r)
    return out, matched, len(master_rows)-matched

# ------------------------ session manifest -------------------------
def build_manifest(master_rows, selection_rows: list[dict[str, Any]] | None = None):
    """Build the Probe 2b manifest.

    Lecanemab uses the exact source-of-truth session whitelist when supplied.
    Other projects retain the automatic date-cluster fallback so the same
    extractor remains usable for Vaccine data.
    """
    probe = [
        r.copy() for r in master_rows
        if txt(r.get("Session_Type_Hint")).lower() == "probe 2b"
        and is_yes(r.get("Include_Primary_Analysis", "YES"))
    ]
    selection_map = {}
    for row in selection_rows or []:
        if not is_yes(row.get("Selected", "YES")):
            continue
        selection_map[selection_key(row)] = row

    by_auto = defaultdict(list)
    by_selected = defaultdict(list)
    for row in probe:
        project = project_text(row)
        if project.casefold() == "lecanemab" and selection_map:
            selected = selection_map.get(selection_key(row))
            if selected:
                rr = row.copy()
                rr["_Source_Timepoint_Months"] = txt(selected.get("Timepoint_Months"))
                rr["_Source_Selection_Order"] = txt(selected.get("Selection_Order"))
                rr["_Source_Chronological_Day"] = txt(selected.get("Source_Chronological_Day"))
                by_selected[(project, animal_text(row), txt(selected.get("Timepoint_Months")))].append(rr)
            continue
        by_auto[(project, animal_text(row))].append(row)

    manifest = []
    qc = []

    # Exact source-of-truth selection path for Lecanemab.
    for (project, animal, tp_text), rows in sorted(by_selected.items()):
        tp = int(float(tp_text)) if tp_text else None
        rows = sorted(rows, key=lambda r: int(float(txt(r.get("_Source_Selection_Order")) or 9999)))
        count_ok = len(rows) == 4
        block_number = {6: 1, 9: 2, 12: 3}.get(tp)
        for si, row in enumerate(rows, 1):
            rr = row.copy()
            rr["Probe_Block_Number"] = block_number or ""
            rr["Timepoint_Months"] = tp or ""
            rr["Timepoint_Label"] = f"{tp} months" if tp else "Unassigned"
            rr["Probe_Session_Number"] = si
            rr["Source_Selection_Order"] = rr.pop("_Source_Selection_Order", "")
            rr["Source_Chronological_Day"] = rr.pop("_Source_Chronological_Day", "")
            design_ok = not (
                project.casefold() == "lecanemab"
                and txt(rr.get("Cohort")).lower().replace(" ", "") in {"cohort1", "1"}
                and tp == 12
            )
            rr["Design_Eligible"] = "YES" if design_ok and count_ok and tp in {6, 9, 12} else "NO"
            if not count_ok:
                rr["Design_Exclusion_Reason"] = f"Source whitelist contains {len(rows)} sessions; exactly 4 required"
            elif not design_ok and tp == 12:
                rr["Design_Exclusion_Reason"] = "Lecanemab Cohort 1 treatment stopped after 9 months"
            elif tp not in {6, 9, 12}:
                rr["Design_Exclusion_Reason"] = "Unassigned timepoint block"
            else:
                rr["Design_Exclusion_Reason"] = ""
            manifest.append(rr)
        qc.append({
            "Project": project,
            "Animal_ID": animal,
            "Probe_Sessions_Final": len(rows),
            "Chronological_Clusters": "SOURCE_WHITELIST",
            "Valid_Clusters_ge4": 1 if count_ok else 0,
            "Assigned_6_9_12": 1 if tp in {6, 9, 12} else 0,
            "Selection_Mode": "SOURCE_WHITELIST",
            "QC_Status": "OK" if count_ok and tp in {6, 9, 12} else "CHECK",
        })

    # Automatic fallback for non-Lecanemab projects.
    by = defaultdict(list)
    for (project, animal), rows in by_auto.items():
        by[(project, animal)].extend(rows)
    for (project, animal), rows in sorted(by.items()):
        clusters = cluster_sessions(rows, 7.0)
        valid_clusters = [c for c in clusters if len(c) >= 4]
        for ci, cluster in enumerate(valid_clusters, 1):
            tp = {1: 6, 2: 9, 3: 12}.get(ci)
            for si, row in enumerate(sorted(cluster, key=lambda z: parse_dt(z.get("Session_DateTime")) or datetime.min), 1):
                rr = row.copy()
                rr["Probe_Block_Number"] = ci
                rr["Timepoint_Months"] = tp or ""
                rr["Timepoint_Label"] = f"{tp} months" if tp else "Unassigned"
                rr["Probe_Session_Number"] = si
                design_ok = not (
                    project.casefold() == "lecanemab"
                    and txt(row.get("Cohort")).lower().replace(" ", "") in {"cohort1", "1"}
                    and tp == 12
                )
                rr["Design_Eligible"] = "YES" if design_ok and tp in {6, 9, 12} and si <= 4 else "NO"
                if not design_ok and tp == 12:
                    rr["Design_Exclusion_Reason"] = "Lecanemab Cohort 1 treatment stopped after 9 months"
                elif si > 4:
                    rr["Design_Exclusion_Reason"] = "Session beyond first 4 valid sessions in timepoint block"
                elif tp not in {6, 9, 12}:
                    rr["Design_Exclusion_Reason"] = "Unassigned timepoint block"
                else:
                    rr["Design_Exclusion_Reason"] = ""
                manifest.append(rr)
        qc.append({
            "Project": project,
            "Animal_ID": animal,
            "Probe_Sessions_Final": len(rows),
            "Chronological_Clusters": "+".join(str(len(c)) for c in clusters),
            "Valid_Clusters_ge4": len(valid_clusters),
            "Assigned_6_9_12": min(len(valid_clusters), 3),
            "Selection_Mode": "AUTOMATIC_DATE_CLUSTER",
            "QC_Status": "OK" if len(valid_clusters) >= 1 else "CHECK",
        })
    return manifest, qc



# ---------------------- matched Stage 4 baseline ----------------------
def build_stage4_manifest(master_rows: list[dict[str,Any]], probe_manifest: list[dict[str,Any]]):
    """Select up to the last four Stage 4 sessions immediately before each Probe block.

    Baseline availability never affects Probe 2b inclusion. For later timepoints,
    Stage 4 sessions must occur after the end of the preceding Probe block so an
    earlier baseline cannot be reused at multiple timepoints.
    """
    stage4_by=defaultdict(list)
    for r in master_rows:
        if txt(r.get("Session_Type_Hint")).lower()=="stage 4" and is_yes(r.get("Include_Primary_Analysis","YES")):
            stage4_by[(txt(r.get("Project_Hint")),txt(r.get("Animal_ID")))].append(r.copy())
    for k in stage4_by:
        stage4_by[k].sort(key=lambda r: parse_dt(r.get("Session_DateTime")) or datetime.min)

    probe_by=defaultdict(list)
    for r in probe_manifest:
        if r.get("Timepoint_Months") in {6,9,12} and int(r.get("Probe_Session_Number") or 0)<=4:
            probe_by[(txt(r.get("Project_Hint")),txt(r.get("Animal_ID")),int(r.get("Probe_Block_Number") or 0))].append(r)

    selected=[]; qc=[]
    animal_blocks=defaultdict(list)
    for (project,animal,block), rows in probe_by.items():
        animal_blocks[(project,animal)].append((block,rows))

    for (project,animal), blocks in sorted(animal_blocks.items()):
        blocks=sorted(blocks,key=lambda x:x[0])
        prev_probe_end=None
        candidates_all=stage4_by.get((project,animal),[])
        for block,prows in blocks:
            prows=sorted(prows,key=lambda r: parse_dt(r.get("Session_DateTime")) or datetime.min)
            first_probe_dt=parse_dt(prows[0].get("Session_DateTime")) if prows else None
            last_probe_dt=parse_dt(prows[-1].get("Session_DateTime")) if prows else None
            tp=prows[0].get("Timepoint_Months") if prows else ""
            candidates=[]
            if first_probe_dt is not None:
                for s4 in candidates_all:
                    dt=parse_dt(s4.get("Session_DateTime"))
                    if dt is None or dt>=first_probe_dt: continue
                    if prev_probe_end is not None and dt<=prev_probe_end: continue
                    candidates.append(s4)
            chosen=candidates[-4:]
            count=len(chosen)
            status="FULL_4" if count>=4 else "PARTIAL_3" if count==3 else "PARTIAL_2" if count==2 else "LIMITED_1" if count==1 else "NO_BASELINE_FOUND"
            for i,s4 in enumerate(chosen,1):
                rr=s4.copy()
                # Metadata/timepoint are inherited from the matched Probe block if absent on Stage 4.
                ref=prows[0]
                for col in ["Cohort","Sex","Genotype","Treatment","Strain"]:
                    if not txt(rr.get(col)): rr[col]=ref.get(col,"")
                rr["Matched_Probe_Block_Number"]=block
                rr["Timepoint_Months"]=tp
                rr["Timepoint_Label"]=f"{tp} months" if tp else ""
                rr["Baseline_Day_Number"]=i
                rr["Baseline_Sessions_Found_In_Window"]=len(candidates)
                rr["Baseline_Sessions_Used"]=count
                rr["Baseline_Status"]=status
                rr["Matched_First_Probe_DateTime"]=txt(prows[0].get("Session_DateTime"))
                rr["Design_Eligible"]="YES"
                rr["Probe_Session_Number"]=i  # temporary compatibility for shared trial enrichment
                selected.append(rr)
            qc.append({
                "Project":project,"Animal_ID":animal,"Cohort":txt(prows[0].get("Cohort")),
                "Sex":txt(prows[0].get("Sex")),"Genotype":txt(prows[0].get("Genotype")),"Treatment":txt(prows[0].get("Treatment")),
                "Timepoint_Months":tp,"Probe_Block_Number":block,
                "First_Probe_DateTime":txt(prows[0].get("Session_DateTime")),
                "Stage4_Found_In_Timepoint_Window":len(candidates),"Stage4_Used":count,
                "Baseline_Status":status,
                "Rule":"Last up to 4 Stage 4 sessions after previous Probe block and before current Probe block"
            })
            if last_probe_dt is not None: prev_probe_end=last_probe_dt
    return selected,qc


def enrich_stage4_trials(raw_trials, stage4_manifest):
    rows=enrich_trials(raw_trials, stage4_manifest)
    lookup={session_key(r):r for r in stage4_manifest}
    for r in rows:
        s=lookup.get((txt(r.get("Database_Fingerprint")),txt(r.get("SID"))),{})
        r["Baseline_Day_Number"]=s.get("Baseline_Day_Number",r.get("Probe_Session_Number"))
        r["Baseline_Status"]=s.get("Baseline_Status","")
        r["Baseline_Sessions_Used"]=s.get("Baseline_Sessions_Used","")
        r["Matched_Probe_Block_Number"]=s.get("Matched_Probe_Block_Number","")
        r["Matched_First_Probe_DateTime"]=s.get("Matched_First_Probe_DateTime","")
        r.pop("Probe_Session_Number",None)
    return rows


def make_stage4_summaries(trials):
    by=defaultdict(list)
    for r in trials: by[(r["Database_Fingerprint"],r["SID"])].append(r)
    session_rows=[]
    for _,g in sorted(by.items(),key=lambda kv:str(kv[0])):
        x=g[0]
        b={k:x.get(k) for k in ["Project","Animal_ID","Cohort","Sex","Genotype","Treatment","Timepoint_Months","Timepoint_Label","Source_Database","SID","Session_DateTime","Baseline_Day_Number","Baseline_Status","Baseline_Sessions_Used","Matched_Probe_Block_Number","Matched_First_Probe_DateTime"]}
        b["Stimulus_Duration_s"]="ALL"
        session_rows.append(summarize_group(g,b))
    av=[]; avgby=defaultdict(list)
    for r in session_rows: avgby[(r["Project"],r["Animal_ID"],r["Timepoint_Months"])].append(r)
    exclude={"Project","Animal_ID","Cohort","Sex","Genotype","Treatment","Timepoint_Months","Timepoint_Label","Source_Database","SID","Session_DateTime","Baseline_Day_Number","Baseline_Status","Baseline_Sessions_Used","Matched_Probe_Block_Number","Matched_First_Probe_DateTime","Stimulus_Duration_s"}
    metrics=[k for k in session_rows[0] if k not in exclude] if session_rows else []
    for _,g in sorted(avgby.items(),key=lambda kv:str(kv[0])):
        b={k:g[0].get(k) for k in ["Project","Animal_ID","Cohort","Sex","Genotype","Treatment","Timepoint_Months","Timepoint_Label","Baseline_Status"]}
        b["Baseline_Sessions_Averaged"]=len(g)
        for c in metrics: b[c]=mean(x.get(c) for x in g)
        av.append(b)
    return session_rows,av


def make_baseline_probe_2s(stage4_avg, probe_avg):
    """Long-format matched 2-s baseline vs Probe table for the baseline-to-probe graphs/models."""
    rows=[]
    for r in stage4_avg:
        rows.append({**r,"Phase":"Stage 4 baseline","Stimulus_Duration_s":2.0})
    for r in probe_avg:
        sd=num(r.get("Stimulus_Duration_s"))
        if sd is None or abs(sd-2.0)>=0.03: continue
        rr=r.copy(); rr["Phase"]="Probe 2b"; rows.append(rr)
    return rows


# -------------------------- trial query ----------------------------
def find_index_db(atlas_dir: Path):
    candidates=[atlas_dir/"cache"/"NMR_ABET_Atlas.sqlite", atlas_dir/"NMR_ABET_Atlas.sqlite"]
    for p in candidates:
        if p.exists(): return p
    hits=list(atlas_dir.rglob("NMR_ABET_Atlas.sqlite"))
    return hits[0] if hits else None

def fetch_trials(db_path: Path, keys: set[tuple[str,str]]):
    conn=sqlite3.connect(str(db_path)); conn.row_factory=sqlite3.Row
    try:
        cols=[r[1] for r in conn.execute("PRAGMA table_info(trials)").fetchall()]
        if not cols: raise RuntimeError("Atlas SQLite has no 'trials' table.")
        out=[]
        # Fetch by fingerprints in batches, filter SID in Python to avoid huge OR expressions.
        fps=sorted({k[0] for k in keys if k[0]})
        for i in range(0,len(fps),500):
            chunk=fps[i:i+500]; ph=",".join("?" for _ in chunk)
            for row in conn.execute(f"SELECT * FROM trials WHERE database_fingerprint IN ({ph})",chunk):
                d=dict(row); k=(txt(d.get("database_fingerprint")),txt(d.get("sid")))
                if k in keys: out.append(d)
        return out
    finally: conn.close()

# ------------------------ trial enrichment -------------------------
def enrich_trials(raw_trials, manifest):
    m={session_key(r):r for r in manifest if txt(r.get("Design_Eligible"))=="YES"}
    rows=[]
    for t in raw_trials:
        k=(txt(t.get("database_fingerprint")),txt(t.get("sid")))
        s=m.get(k)
        if not s: continue
        out={
            "Project":txt(s.get("Project_Hint")),"Animal_ID":txt(s.get("Animal_ID")),"Cohort":txt(s.get("Cohort")),
            "Sex":txt(s.get("Sex")),"Genotype":txt(s.get("Genotype")),"Treatment":txt(s.get("Treatment")),
            "Timepoint_Months":s.get("Timepoint_Months"),"Timepoint_Label":s.get("Timepoint_Label"),
            "Probe_Session_Number":s.get("Probe_Session_Number"),"Source_Database":txt(s.get("Source_Database")),
            "Database_Fingerprint":txt(s.get("Database_Fingerprint")),"SID":txt(s.get("SID")),"Session_DateTime":txt(s.get("Session_DateTime")),
            "Schedule_Name":txt(s.get("Schedule_Name")),"Trial_Index":t.get("trial_index"),"Trial_Counter":t.get("trial_counter"),
            "Stimulus_Onset":t.get("stimulus_onset"),"Outcome_Time":t.get("outcome_time"),"Outcome":txt(t.get("outcome")),
            "Stimulus_Duration_s":norm_sd(t.get("stimulus_duration_s")),"Current_Image":t.get("current_image"),"Correct_Image":t.get("correct_image"),
            "Current_ITI_s":t.get("current_iti_s"),"Correct_Grid_Position":t.get("correct_grid_position"),
            "Hit":int(t.get("hit") or 0),"Miss":int(t.get("miss") or 0),"False_Alarm":int(t.get("false_alarm") or 0),"Correct_Rejection":int(t.get("correct_rejection") or 0),
            "Correction_Trial_Correct_Rejection":int(t.get("correction_trial_correct_rejection") or 0),"Correction_Trial_Mistake":int(t.get("correction_trial_mistake") or 0),
            "Response_Latency_Raw_s":num(t.get("response_latency_s")),"Reward_Retrieval_Latency_Raw_s":num(t.get("reward_retrieval_latency_s")),
        }
        out["Correct_Touch_Latency_Raw_s"] = out["Response_Latency_Raw_s"] if out["Hit"] else None
        out["Incorrect_Touch_Latency_Raw_s"] = out["Response_Latency_Raw_s"] if out["False_Alarm"] else None
        rows.append(out)
    # session-level IQR screening separately by class, min 4 values; raw retained.
    by=defaultdict(list)
    for i,r in enumerate(rows): by[(r["Database_Fingerprint"],r["SID"])].append((i,r))
    for _,items in by.items():
        for rawcol,clean,colflag,prefix in [
            ("Correct_Touch_Latency_Raw_s","Correct_Touch_Latency_IQRclean_s","Correct_Touch_Latency_IQR_Excluded","Correct"),
            ("Incorrect_Touch_Latency_Raw_s","Incorrect_Touch_Latency_IQRclean_s","Incorrect_Touch_Latency_IQR_Excluded","Incorrect"),
            ("Reward_Retrieval_Latency_Raw_s","Reward_Retrieval_Latency_IQRclean_s","Reward_Retrieval_Latency_IQR_Excluded","Reward"),
        ]:
            vals=[r[rawcol] for _,r in items if r[rawcol] is not None]
            lo,hi=iqr_bounds(vals)
            for _,r in items:
                v=r[rawcol]; excl = bool(v is not None and lo is not None and (v<lo or v>hi))
                r[clean]=None if excl else v; r[colflag]=1 if excl else 0; r[f"{prefix}_IQR_Lower_s"]=lo; r[f"{prefix}_IQR_Upper_s"]=hi
        n=len(items)
        ordered=sorted(items,key=lambda x:(num(x[1].get("Trial_Index")) or 0))
        for rank,(_,r) in enumerate(ordered,1):
            r["Session_Progress_Decile"]=min(10,max(1,math.ceil(rank/n*10))) if n else None
            r["Session_Progress_Percent_Midpoint"]=r["Session_Progress_Decile"]*10-5 if n else None
        # next-trial previous outcome within each session
        prev=""
        for _,r in ordered:
            r["Previous_Trial_Outcome"]=prev
            if prev in {"Hit","Correct Rejection"}: r["Previous_Outcome_Class"]="After correct trial"
            elif prev in {"Miss","False Alarm","Mistake"}: r["Previous_Outcome_Class"]="After error/miss"
            else: r["Previous_Outcome_Class"]=""
            prev=r["Outcome"]
    return rows

# --------------------------- summaries -----------------------------
def summarize_group(group: list[dict[str,Any]], base: dict[str,Any]):
    h=sum(r["Hit"] for r in group); m=sum(r["Miss"] for r in group); fa=sum(r["False_Alarm"] for r in group); cr=sum(r["Correct_Rejection"] for r in group)
    out=base.copy(); out.update({"Hits":h,"Misses":m,"False_Alarms":fa,"Correct_Rejections":cr,"Signal_Trials":h+m,"Noise_Trials":fa+cr,"Core_Trials":h+m+fa+cr})
    out.update(rate_metrics(h,m,fa,cr))
    out.update({
        "Mean_Correct_Touch_Latency_Raw_s":mean(r["Correct_Touch_Latency_Raw_s"] for r in group),
        "Mean_Correct_Touch_Latency_IQRclean_s":mean(r["Correct_Touch_Latency_IQRclean_s"] for r in group),
        "Correct_Latency_IQR_Excluded_n":sum(r["Correct_Touch_Latency_IQR_Excluded"] for r in group),
        "Mean_Incorrect_Touch_Latency_Raw_s":mean(r["Incorrect_Touch_Latency_Raw_s"] for r in group),
        "Mean_Incorrect_Touch_Latency_IQRclean_s":mean(r["Incorrect_Touch_Latency_IQRclean_s"] for r in group),
        "Incorrect_Latency_IQR_Excluded_n":sum(r["Incorrect_Touch_Latency_IQR_Excluded"] for r in group),
        "Mean_Reward_Retrieval_Latency_Raw_s":mean(r["Reward_Retrieval_Latency_Raw_s"] for r in group),
        "Mean_Reward_Retrieval_Latency_IQRclean_s":mean(r["Reward_Retrieval_Latency_IQRclean_s"] for r in group),
        "Reward_Latency_IQR_Excluded_n":sum(r["Reward_Retrieval_Latency_IQR_Excluded"] for r in group),
    })
    return out

def base_fields(r):
    return {k:r.get(k) for k in ["Project","Animal_ID","Cohort","Sex","Genotype","Treatment","Timepoint_Months","Timepoint_Label","Probe_Session_Number","Source_Database","SID","Session_DateTime"]}

def make_summaries(trials):
    session_sd=[]; by=defaultdict(list)
    for r in trials: by[(r["Database_Fingerprint"],r["SID"],r["Stimulus_Duration_s"])].append(r)
    for _,g in sorted(by.items(),key=lambda kv:str(kv[0])):
        b=base_fields(g[0]); b["Stimulus_Duration_s"]=g[0]["Stimulus_Duration_s"]; session_sd.append(summarize_group(g,b))

    dec=[]; by=defaultdict(list)
    for r in trials: by[(r["Database_Fingerprint"],r["SID"],r["Session_Progress_Decile"],r["Stimulus_Duration_s"])].append(r)
    for _,g in sorted(by.items(),key=lambda kv:str(kv[0])):
        b=base_fields(g[0]); b["Session_Progress_Decile"]=g[0]["Session_Progress_Decile"]; b["Session_Progress_Percent_Midpoint"]=g[0]["Session_Progress_Percent_Midpoint"]; b["Stimulus_Duration_s"]=g[0]["Stimulus_Duration_s"]; dec.append(summarize_group(g,b))

    # Four-session animal/timepoint averages: average session-level metrics, preserving animal as unit.
    metric_cols=[k for k in session_sd[0].keys() if k not in {"Project","Animal_ID","Cohort","Sex","Genotype","Treatment","Timepoint_Months","Timepoint_Label","Probe_Session_Number","Source_Database","SID","Session_DateTime","Stimulus_Duration_s"}] if session_sd else []
    av=[]; by=defaultdict(list)
    for r in session_sd: by[(r["Project"],r["Animal_ID"],r["Timepoint_Months"],r["Stimulus_Duration_s"])].append(r)
    for _,g in sorted(by.items(),key=lambda kv:str(kv[0])):
        b={k:g[0].get(k) for k in ["Project","Animal_ID","Cohort","Sex","Genotype","Treatment","Timepoint_Months","Timepoint_Label","Stimulus_Duration_s"]}
        b["Sessions_Averaged"]=len({x["Probe_Session_Number"] for x in g})
        for c in metric_cols: b[c]=mean(x.get(c) for x in g)
        av.append(b)

    # Next-trial long rows (trial-level outcomes retained for flexible models)
    nxt=[]
    for r in trials:
        if not r["Previous_Outcome_Class"]: continue
        nxt.append({
            **{k:r.get(k) for k in ["Project","Animal_ID","Cohort","Sex","Genotype","Treatment","Timepoint_Months","Timepoint_Label","Probe_Session_Number","Source_Database","SID","Trial_Index","Stimulus_Duration_s"]},
            "Previous_Trial_Outcome":r["Previous_Trial_Outcome"],"Previous_Outcome_Class":r["Previous_Outcome_Class"],"Current_Outcome":r["Outcome"],
            "Current_Hit":r["Hit"],"Current_Miss":r["Miss"],"Current_False_Alarm":r["False_Alarm"],"Current_Correct_Rejection":r["Correct_Rejection"],
            "Correct_Touch_Latency_IQRclean_s":r["Correct_Touch_Latency_IQRclean_s"],"Incorrect_Touch_Latency_IQRclean_s":r["Incorrect_Touch_Latency_IQRclean_s"],"Reward_Retrieval_Latency_IQRclean_s":r["Reward_Retrieval_Latency_IQRclean_s"]
        })
    return session_sd,av,dec,nxt



def make_taylor49(trials: list[dict[str,Any]], manifest: list[dict[str,Any]]):
    """Build the exact legacy 49-column Taylor-compatible session summary.

    One row is written per design-eligible QC-approved Probe 2b session.
    Latency values use the same session-level 1.5x IQR-cleaned trial fields as
    the main extractor outputs. AVG_count includes all reconstructed image
    presentations assigned to the stimulus duration, matching the old converter.
    """
    tby=defaultdict(list)
    for r in trials:
        tby[(txt(r.get("Database_Fingerprint")),txt(r.get("SID")))].append(r)
    rows=[]
    for s in sorted([r for r in manifest if txt(r.get("Design_Eligible"))=="YES"],
                    key=lambda r:(txt(r.get("Project_Hint")),txt(r.get("Animal_ID")),parse_dt(r.get("Session_DateTime")) or datetime.min)):
        key=session_key(s); g=tby.get(key,[])
        row={c:"" for c in TAYLOR_49_COLUMNS}
        row.update({
            "AnimalID":txt(s.get("Animal_ID")),
            "Age":s.get("Timepoint_Months", ""),
            "Sex":txt(s.get("Sex")),
            "Treatment":txt(s.get("Treatment")),
            "Genotype":txt(s.get("Genotype")),
            "Strain":txt(s.get("Strain")) or txt(s.get("Genotype")),
            "ExpName":txt(s.get("Project_Hint")),
            "SessionName":txt(s.get("Schedule_Name")),
            "sessionInfoName":txt(s.get("SID")),
            "S_ID":txt(s.get("SID")),
        })
        row["AVG_End Summary - Hits"] = sum(int(r.get("Hit") or 0) for r in g)
        for dur,label in TAYLOR_DURATIONS:
            dg=[r for r in g if r.get("Stimulus_Duration_s") is not None and abs(float(r["Stimulus_Duration_s"])-dur)<0.03]
            if not dg: continue
            row[f"AVG_Average - Correct Choice Latency at {label} SD"] = mean(r.get("Correct_Touch_Latency_IQRclean_s") for r in dg if int(r.get("Hit") or 0))
            row[f"AVG_Average - Mistake Latency at {label} SD"] = mean(r.get("Incorrect_Touch_Latency_IQRclean_s") for r in dg if int(r.get("False_Alarm") or 0))
            row[f"AVG_Average - Reward Retrieval Latency at {label} SD"] = mean(r.get("Reward_Retrieval_Latency_IQRclean_s") for r in dg if int(r.get("Hit") or 0))
            row[f"AVG_count at {label}"] = len(dg)
            row[f"AVG_End Summary - CorrectRejection at {label} SD"] = sum(int(r.get("Correct_Rejection") or 0) for r in dg)
            row[f"AVG_End Summary - Hits at {label} SD"] = sum(int(r.get("Hit") or 0) for r in dg)
            row[f"AVG_End Summary - Miss at {label} SD"] = sum(int(r.get("Miss") or 0) for r in dg)
            row[f"AVG_End Summary - Mistake at {label} SD"] = sum(int(r.get("False_Alarm") or 0) for r in dg)
        rows.append(row)
    return rows


def make_legacy67(trials: list[dict[str,Any]], manifest: list[dict[str,Any]]):
    """Build the exact legacy 67-column extended CPT session summary."""
    base_rows=make_taylor49(trials, manifest)
    eligible=sorted([r for r in manifest if txt(r.get("Design_Eligible"))=="YES"],
                    key=lambda r:(txt(r.get("Project_Hint")),txt(r.get("Animal_ID")),parse_dt(r.get("Session_DateTime")) or datetime.min))
    tby=defaultdict(list)
    for r in trials:
        tby[(txt(r.get("Database_Fingerprint")),txt(r.get("SID")))].append(r)
    out=[]
    for s,base in zip(eligible,base_rows):
        row={c:"" for c in LEGACY_67_COLUMNS}
        row.update(base)
        row["Date_Time"]=txt(s.get("Session_DateTime"))
        row["chronological day"]=s.get("Probe_Session_Number","")
        g=tby.get(session_key(s),[])
        for dur,label in TAYLOR_DURATIONS:
            dg=[r for r in g if r.get("Stimulus_Duration_s") is not None and abs(float(r["Stimulus_Duration_s"])-dur)<0.03]
            h=sum(int(r.get("Hit") or 0) for r in dg)
            m=sum(int(r.get("Miss") or 0) for r in dg)
            fa=sum(int(r.get("False_Alarm") or 0) for r in dg)
            cr=sum(int(r.get("Correct_Rejection") or 0) for r in dg)
            met=rate_metrics(h,m,fa,cr)
            row[f"Hit Rate at {label} SD"]=met["Hit_Rate"] if met["Hit_Rate"] is not None else ""
            row[f"False Alarm Rate at {label} SD"]=met["False_Alarm_Rate"] if met["False_Alarm_Rate"] is not None else ""
            row[f"dPrime at {label} SD"]=met["d_prime"] if met["d_prime"] is not None else ""
            row[f"c Response Bias at {label} SD"]=met["criterion_c"] if met["criterion_c"] is not None else ""
        out.append(row)
    return out

# ----------------------------- main --------------------------------
def run(atlas_dir: Path, output_dir: Path | None=None):
    reports=atlas_dir/"reports" if (atlas_dir/"reports").exists() else atlas_dir
    master_candidates=[
        reports/"ATLAS_SESSION_MASTER_FINAL.csv",
        atlas_dir/"ATLAS_SESSION_MASTER_FINAL.csv",
        reports/"ATLAS_SESSION_MASTER.csv",
        atlas_dir/"ATLAS_SESSION_MASTER.csv",
    ]
    master=next((p for p in master_candidates if p.exists()), None)
    if not master:
        searched="\n".join(str(p) for p in master_candidates)
        raise FileNotFoundError(f"Could not find ATLAS_SESSION_MASTER_FINAL.csv or ATLAS_SESSION_MASTER.csv. Searched:\n{searched}")
    db=find_index_db(atlas_dir)
    if not db: raise FileNotFoundError("Could not find cache/NMR_ABET_Atlas.sqlite under the selected Atlas folder")
    out=output_dir or (atlas_dir/"probe2b_analysis")
    out.mkdir(parents=True,exist_ok=True)

    master_rows=read_csv(master)
    metadata_path=find_metadata_csv(reports) or find_metadata_csv(atlas_dir)
    if not metadata_path:
        raise FileNotFoundError(f"Could not find ATLAS_ANIMAL_METADATA.csv beside {master}. The extractor needs finalized Project + Animal ID metadata for Sex/Genotype/Treatment/Cohort.")
    metadata_rows=read_csv(metadata_path)
    master_rows, metadata_matched, metadata_unmatched = merge_animal_metadata(master_rows, metadata_rows)

    reconciliation_path = find_reconciliation_workbook(atlas_dir, reports)
    reconciliation_rows = []
    reconciliation_audit = []
    reconciliation_counts = {}
    if reconciliation_path:
        reconciliation_rows = read_decision_workbook(reconciliation_path)
        master_rows, reconciliation_audit, reconciliation_counts = apply_reconciliation_decisions(master_rows, reconciliation_rows)
        write_csv(out/"ATLAS_RECONCILIATION_DECISIONS_APPLIED.csv", reconciliation_audit)
    else:
        write_csv(out/"ATLAS_RECONCILIATION_DECISIONS_APPLIED.csv", [])

    # Finalized ATLAS_ANIMAL_METADATA is the approved animal population. Sessions
    # for animals absent from that table remain in Atlas for audit, but do not
    # enter Probe 2b / Stage 4 analysis exports.
    analysis_master_rows=[
        r for r in master_rows
        if txt(r.get("Metadata_Match_Status")) in {"MATCHED", "MATCHED_BY_RECONCILIATION"}
        and txt(r.get("Reconciliation_Include") or "YES") == "YES"
    ]
    selection_path=find_selection_csv(reports) or find_selection_csv(atlas_dir)
    selection_rows=read_csv(selection_path) if selection_path else []
    if selection_rows and any(txt(r.get("Project_Group") or r.get("Project_Hint") or r.get("Project")).casefold()=="lecanemab" for r in selection_rows):
        # Corrected Sheet2 S_ID values are not Atlas session IDs. The mapped
        # whitelist must carry the Atlas database fingerprint and SID. Refuse
        # the older unmapped file rather than silently matching coincidental
        # numeric IDs.
        required = {"Atlas_Master_SID", "Atlas_Master_Database_Fingerprint"}
        missing = sorted(required - set(selection_rows[0].keys()))
        if missing:
            raise RuntimeError(
                "The Lecanemab whitelist is the older unmapped file. Replace "
                "reports\\ATLAS_PROBE2B_SESSION_SELECTIONS.csv with the corrected "
                "package version; it must contain Atlas_Master_SID and "
                "Atlas_Master_Database_Fingerprint."
            )
    if selection_rows:
        master_probe_keys = defaultdict(int)
        for row in analysis_master_rows:
            if txt(row.get("Session_Type_Hint")).lower() == "probe 2b":
                master_probe_keys[selection_key(row)] += 1
        selection_audit=[]
        for row in selection_rows:
            key=selection_key(row)
            audit=row.copy()
            audit["Atlas_Master_Match"]="MATCHED" if master_probe_keys.get(key,0) else "MISSING_IN_ATLAS_MASTER"
            audit["Atlas_Master_Match_Count"]=str(master_probe_keys.get(key,0))
            selection_audit.append(audit)
        write_csv(out/"PROBE2B_SESSION_SELECTION_AUDIT.csv",selection_audit)
    manifest,qc=build_manifest(analysis_master_rows, selection_rows)
    write_csv(out/"PROBE2B_SESSION_MANIFEST.csv",manifest)
    write_csv(out/"PROBE2B_TIMEPOINT_ASSIGNMENT_QC.csv",qc)
    eligible=[r for r in manifest if txt(r.get("Design_Eligible"))=="YES"]
    keys={session_key(r) for r in eligible}
    raw=fetch_trials(db,keys)
    trials=enrich_trials(raw,eligible)
    if not trials: raise RuntimeError("No matching trial rows were found for the eligible Probe 2b sessions. Check Atlas cache and final master paths.")
    write_csv(out/"PROBE2B_TRIAL_MASTER.csv",trials)
    session_sd,av,dec,nxt=make_summaries(trials)
    write_csv(out/"PROBE2B_SESSION_SD_METRICS.csv",session_sd)
    write_csv(out/"PROBE2B_TIMEPOINT_AVERAGES.csv",av)
    write_csv(out/"PROBE2B_WITHIN_SESSION_DECILES.csv",dec)
    write_csv(out/"PROBE2B_NEXT_TRIAL.csv",nxt)
    taylor49=make_taylor49(trials,eligible)
    write_csv(out/"CPT_49_COLUMN_ALL_DATABASES.csv",taylor49,TAYLOR_49_COLUMNS)
    legacy67=make_legacy67(trials,eligible)
    write_csv(out/"CPT_67_COLUMN_ALL_DATABASES.csv",legacy67,LEGACY_67_COLUMNS)

    # Optional matched Stage 4 baseline layer. Missing/partial Stage 4 never excludes Probe 2b.
    stage4_manifest,stage4_qc=build_stage4_manifest(analysis_master_rows,manifest)
    write_csv(out/"STAGE4_BASELINE_SESSION_MANIFEST.csv",stage4_manifest)
    write_csv(out/"STAGE4_BASELINE_QC.csv",stage4_qc)
    stage4_trials=[]; stage4_session=[]; stage4_avg=[]; baseline_probe=[]
    if stage4_manifest:
        stage4_keys={session_key(r) for r in stage4_manifest}
        stage4_raw=fetch_trials(db,stage4_keys)
        stage4_trials=enrich_stage4_trials(stage4_raw,stage4_manifest)
        write_csv(out/"STAGE4_BASELINE_TRIAL_MASTER.csv",stage4_trials)
        stage4_session,stage4_avg=make_stage4_summaries(stage4_trials)
        write_csv(out/"STAGE4_BASELINE_SESSION_METRICS.csv",stage4_session)
        write_csv(out/"STAGE4_BASELINE_TIMEPOINT_AVERAGES.csv",stage4_avg)
        baseline_probe=make_baseline_probe_2s(stage4_avg,av)
        write_csv(out/"BASELINE_TO_PROBE_2S_LONG.csv",baseline_probe)
    else:
        write_csv(out/"STAGE4_BASELINE_TRIAL_MASTER.csv",[])
        write_csv(out/"STAGE4_BASELINE_SESSION_METRICS.csv",[])
        write_csv(out/"STAGE4_BASELINE_TIMEPOINT_AVERAGES.csv",[])
        write_csv(out/"BASELINE_TO_PROBE_2S_LONG.csv",[])

    sessions_expected=len(keys); sessions_found=len({(r["Database_Fingerprint"],r["SID"]) for r in trials})
    summary=[
        {"Item":"Atlas final master","Value":str(master)},
        {"Item":"Atlas trial cache","Value":str(db)},
        {"Item":"Atlas animal metadata","Value":str(metadata_path)},
        {"Item":"Atlas reconciliation decision workbook","Value":str(reconciliation_path) if reconciliation_path else "NOT FOUND — metadata-only filtering used"},
        {"Item":"Reconciliation decision rows read","Value":len(reconciliation_rows)},
        {"Item":"Reconciliation audit rows written","Value":len(reconciliation_audit)},
        {"Item":"Reconciliation KEEP session rows","Value":reconciliation_counts.get("KEEP", 0)},
        {"Item":"Reconciliation EXCLUDE session rows","Value":reconciliation_counts.get("EXCLUDE", 0)},
        {"Item":"Reconciliation REMAP session rows","Value":reconciliation_counts.get("REMAP", 0)},
        {"Item":"Reconciliation unresolved session rows excluded","Value":reconciliation_counts.get("UNRESOLVED", 0) + reconciliation_counts.get("REMAP_MISSING_CORRECT_ID", 0)},
        {"Item":"Exact Probe 2b session whitelist","Value":str(selection_path) if selection_path else "NOT FOUND — automatic session selection used"},
        {"Item":"Exact whitelist rows","Value":len(selection_rows)},
        {"Item":"Session-master rows matched to Project + Animal ID metadata","Value":metadata_matched},
        {"Item":"Session-master rows unmatched to Project + Animal ID metadata (retained in Atlas, excluded from analysis exports)","Value":metadata_unmatched},
        {"Item":"Approved Project + Animal combinations in finalized metadata","Value":len({(txt(r.get("Project_Group") or r.get("Project_Hint")),txt(r.get("AnimalID") or r.get("Animal_ID"))) for r in metadata_rows if txt(r.get("AnimalID") or r.get("Animal_ID"))})},
        {"Item":"Probe sessions in QC-approved manifest","Value":len(manifest)},
        {"Item":"Probe sessions eligible after design rules","Value":sessions_expected},
        {"Item":"Eligible sessions with cached trial data","Value":sessions_found},
        {"Item":"Extracted trial rows","Value":len(trials)},
        {"Item":"Session × stimulus-duration rows","Value":len(session_sd)},
        {"Item":"Animal × timepoint × stimulus-duration average rows","Value":len(av)},
        {"Item":"Within-session decile rows","Value":len(dec)},
        {"Item":"Next-trial rows","Value":len(nxt)},
        {"Item":"Taylor-compatible 49-column session rows","Value":len(taylor49)},
        {"Item":"Legacy extended 67-column session rows","Value":len(legacy67)},
        {"Item":"Matched Stage 4 baseline sessions selected","Value":len(stage4_manifest)},
        {"Item":"Matched Stage 4 trial rows","Value":len(stage4_trials)},
        {"Item":"Animal × timepoint Stage 4 baseline averages","Value":len(stage4_avg)},
        {"Item":"Baseline-to-Probe 2s long rows","Value":len(baseline_probe)},
        {"Item":"Timepoints with full 4-session Stage 4 baseline","Value":sum(1 for r in stage4_qc if r.get("Baseline_Status")=="FULL_4")},
        {"Item":"Timepoints with partial/missing Stage 4 baseline","Value":sum(1 for r in stage4_qc if r.get("Baseline_Status")!="FULL_4")},
    ]
    write_csv(out/"PROBE2B_EXTRACTION_SUMMARY.csv",summary,["Item","Value"])
    return {"output":out,"summary":summary,"sessions_expected":sessions_expected,"sessions_found":sessions_found,"trials":len(trials)}

def gui():
    import tkinter as tk
    from tkinter import filedialog,messagebox,ttk
    root=tk.Tk(); root.title("Atlas Probe 2b Extractor"); root.geometry("800x560")
    frm=ttk.Frame(root,padding=14); frm.pack(fill="both",expand=True)
    atlas=tk.StringVar(); out=tk.StringVar()
    ttk.Label(frm,text="NMR ABET Atlas — Probe 2b + Stage 4 Extractor",font=("Segoe UI",18,"bold")).grid(row=0,column=0,columnspan=3,sticky="w")
    ttk.Label(frm,text="Uses the QC-final Probe session master + Atlas trial cache, and optionally matches up to the last 4 Stage 4 baselines before each Probe block. It does not change Atlas or the raw databases.").grid(row=1,column=0,columnspan=3,sticky="w",pady=(0,14))
    def bd(var):
        p=filedialog.askdirectory();
        if p: var.set(p)
    for row,label,var in [(2,"Atlas output folder",atlas),(3,"Output folder (optional)",out)]:
        ttk.Label(frm,text=label).grid(row=row,column=0,sticky="w",pady=5)
        ttk.Entry(frm,textvariable=var,width=70).grid(row=row,column=1,sticky="ew",padx=8)
        ttk.Button(frm,text="Browse…",command=lambda v=var:bd(v)).grid(row=row,column=2)
    text=tk.Text(frm,height=20,wrap="word",font=("Consolas",9)); text.grid(row=5,column=0,columnspan=3,sticky="nsew",pady=(12,0))
    frm.columnconfigure(1,weight=1); frm.rowconfigure(5,weight=1)
    def go():
        if not atlas.get(): messagebox.showerror("Missing Atlas folder","Choose the Atlas output folder."); return
        text.delete("1.0","end")
        try:
            text.insert("end","Reading QC-approved Probe sessions, Atlas trial cache, and matched Stage 4 baselines…\n"); root.update_idletasks()
            res=run(Path(atlas.get()),Path(out.get()) if out.get() else None)
            for x in res["summary"]: text.insert("end",f"{x['Item']}: {x['Value']}\n")
            text.insert("end",f"\nDONE\nOutputs: {res['output']}\n")
            messagebox.showinfo("Probe 2b extraction complete",f"Done.\n\nSessions matched: {res['sessions_found']}/{res['sessions_expected']}\nTrials: {res['trials']:,}\n\nOutputs:\n{res['output']}")
        except Exception as e:
            text.insert("end","\n"+traceback.format_exc()); messagebox.showerror("Extractor error",str(e))
    ttk.Button(frm,text="EXTRACT QC'D PROBE 2B + STAGE 4 DATA",command=go).grid(row=6,column=0,columnspan=3,sticky="ew",pady=(10,0),ipady=7)
    root.mainloop()

def main():
    p=argparse.ArgumentParser(description="Extract QC-approved Probe 2b trials, CPT metrics, and matched Stage 4 baselines from NMR ABET Atlas")
    p.add_argument("--atlas",help="Atlas output folder containing reports/ and cache/")
    p.add_argument("--output",help="Optional output folder")
    a=p.parse_args()
    if a.atlas:
        r=run(Path(a.atlas),Path(a.output) if a.output else None)
        print("DONE",r["output"]); return 0
    gui(); return 0

if __name__=="__main__":
    raise SystemExit(main())
