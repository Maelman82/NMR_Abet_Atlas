from __future__ import annotations

import csv
from collections import defaultdict
from pathlib import Path
from typing import Any

from .common import norm_text

CANON = {
    "animalid": "AnimalID",
    "animal id": "AnimalID",
    "project_group": "Project_Group",
    "project group": "Project_Group",
    "project": "Project_Group",
    "cohort": "Cohort",
    "timepoint_label": "Timepoint_Label",
    "timepoint label": "Timepoint_Label",
    "sex": "Sex",
    "genotype": "Genotype",
    "treatment": "Treatment",
    "strain": "Strain",
}

DECISION_CANON = {
    "atlas_db_animal_id": "Atlas_DB_Animal_ID",
    "atlas db animal id": "Atlas_DB_Animal_ID",
    "animal_id": "Atlas_DB_Animal_ID",
    "animal id": "Atlas_DB_Animal_ID",
    "atlas_project_hint": "Atlas_Project_Hint",
    "atlas project hint": "Atlas_Project_Hint",
    "project_hint": "Atlas_Project_Hint",
    "project hint": "Atlas_Project_Hint",
    "decision": "Decision",
    "correct_animal_id": "Correct_Animal_ID",
    "correct animal id": "Correct_Animal_ID",
    "final_project_group": "Final_Project_Group",
    "final project group": "Final_Project_Group",
    "project_group": "Final_Project_Group",
    "project group": "Final_Project_Group",
    "final_cohort": "Final_Cohort",
    "final cohort": "Final_Cohort",
    "cohort": "Final_Cohort",
    "final_sex": "Final_Sex",
    "final sex": "Final_Sex",
    "sex": "Final_Sex",
    "final_genotype": "Final_Genotype",
    "final genotype": "Final_Genotype",
    "genotype": "Final_Genotype",
    "final_treatment": "Final_Treatment",
    "final treatment": "Final_Treatment",
    "treatment": "Final_Treatment",
    "reviewer_notes": "Reviewer_Notes",
    "reviewer notes": "Reviewer_Notes",
    "notes": "Reviewer_Notes",
}


def load_metadata_csv(path: Path | None) -> list[dict[str, str]]:
    if path is None:
        return []
    rows: list[dict[str, str]] = []
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        r = csv.DictReader(f)
        for raw in r:
            row: dict[str, str] = {}
            for k, v in raw.items():
                if not k:
                    continue
                canon = CANON.get(k.strip().lower(), k.strip())
                row[canon] = norm_text(v)
            if row.get("AnimalID"):
                rows.append(row)
    return rows


def _canon_decision_row(raw: dict[str, Any]) -> dict[str, str]:
    row: dict[str, str] = {}
    for k, v in raw.items():
        if not k:
            continue
        canon = DECISION_CANON.get(str(k).strip().lower(), str(k).strip())
        row[canon] = norm_text(v)
    return row


def load_reconciliation_decisions(path: Path | None) -> list[dict[str, str]]:
    """Load the human-reviewed Atlas reconciliation decision register.

    CSV is supported directly. XLSX is supported when openpyxl is installed; the
    expected sheet is named "Decisions", with a fallback to the active sheet.
    """
    if path is None:
        return []
    path = path.expanduser()
    if not path.exists():
        return []

    raw_rows: list[dict[str, Any]] = []
    if path.suffix.lower() in {".csv", ".txt"}:
        with path.open("r", encoding="utf-8-sig", newline="") as f:
            raw_rows = list(csv.DictReader(f))
    elif path.suffix.lower() in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError(
                "ATLAS_RECONCILIATION_DECISIONS.xlsx needs openpyxl. "
                "Run: py -3 -m pip install openpyxl"
            ) from exc
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb["Decisions"] if "Decisions" in wb.sheetnames else wb.active
        rows = list(ws.iter_rows(values_only=True))
        if rows:
            headers = [norm_text(h) for h in rows[0]]
            for vals in rows[1:]:
                raw_rows.append({headers[i]: vals[i] if i < len(vals) else "" for i in range(len(headers)) if headers[i]})
    else:
        raise RuntimeError(f"Unsupported reconciliation decision file type: {path.suffix}")

    decisions: list[dict[str, str]] = []
    for raw in raw_rows:
        row = _canon_decision_row(raw)
        animal = row.get("Atlas_DB_Animal_ID", "")
        if not animal:
            continue
        decision = row.get("Decision", "").strip().upper()
        if decision in {"CORRECT", "CORRECT ID"}:
            decision = "REMAP"
        if decision not in {"", "KEEP", "EXCLUDE", "REMAP", "UNSURE"}:
            decision = "UNSURE"
        row["Decision"] = decision
        decisions.append(row)
    return decisions


def animal_metadata_summary(rows: list[dict[str, str]]) -> list[dict[str, str]]:
    """Collapse repeated timepoint rows while explicitly flagging conflicts."""
    groups: dict[tuple[str, str], list[dict[str, str]]] = defaultdict(list)
    for row in rows:
        groups[(row.get("Project_Group", ""), row.get("AnimalID", ""))].append(row)

    out: list[dict[str, str]] = []
    for (project, animal), members in sorted(groups.items()):
        rec = {"Project_Group": project, "AnimalID": animal}
        conflicts: list[str] = []
        for field in ("Cohort", "Sex", "Genotype", "Treatment", "Strain"):
            vals = sorted({norm_text(m.get(field)) for m in members if norm_text(m.get(field))})
            rec[field] = vals[0] if len(vals) == 1 else " | ".join(vals)
            if len(vals) > 1:
                conflicts.append(field)
        tps = sorted({norm_text(m.get("Timepoint_Label")) for m in members if norm_text(m.get("Timepoint_Label"))})
        rec["Timepoints_In_Metadata"] = " | ".join(tps)
        rec["Metadata_Status"] = "CONFLICT" if conflicts else "OK"
        rec["Metadata_Conflict_Fields"] = "; ".join(conflicts)
        out.append(rec)
    return out
